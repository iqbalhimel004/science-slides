from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import subprocess

base = Path('/mnt/data/science_slides_v21')
v02_render_dir = base / 'full_lesson2_v02' / 'Lesson2_Controlled_Full_v0_2'
v04_dir = base / 'full_lesson2_v04'
out = base / 'full_lesson2_v05'
out.mkdir(parents=True, exist_ok=True)

# Rendered deck images are 1600 x 900 for a 13.333 x 7.5 inch deck.
SCALE = 120

crops = {
    'fish_left_panel.png': (
        v02_render_dir / 'slide-9.png',
        (0.65*SCALE, 1.95*SCALE, (0.65+8.25)*SCALE, (1.95+4.75)*SCALE),
    ),
    'mirage_left_panel.png': (
        v02_render_dir / 'slide-10.png',
        (0.65*SCALE, 1.95*SCALE, (0.65+5.15)*SCALE, (1.95+4.75)*SCALE),
    ),
    'mirage_right_panel.png': (
        v02_render_dir / 'slide-10.png',
        (6.05*SCALE, 1.95*SCALE, (6.05+6.62)*SCALE, (1.95+4.75)*SCALE),
    ),
}

for name, (src, box) in crops.items():
    image = Image.open(src).convert('RGB')
    image.crop(tuple(map(int, box))).save(out / name, quality=95)

prs = Presentation(v04_dir / 'Lesson2_Controlled_Full_v0_4.pptx')

# User explicitly requested image-only restoration. Do not change text or any non-image slide content.
# Slide indices are 0-based: slide 9 = fish, slide 10 = mirage.
prs.slides[8].shapes.add_picture(str(out / 'fish_left_panel.png'), Inches(0.65), Inches(1.95), Inches(8.25), Inches(4.75))
prs.slides[9].shapes.add_picture(str(out / 'mirage_left_panel.png'), Inches(0.65), Inches(1.95), Inches(5.15), Inches(4.75))
prs.slides[9].shapes.add_picture(str(out / 'mirage_right_panel.png'), Inches(6.05), Inches(1.95), Inches(6.62), Inches(4.75))

pptx_out = out / 'Lesson2_Controlled_Full_v0_5.pptx'
prs.save(pptx_out)

subprocess.run(['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(pptx_out)], check=True)
subprocess.run(['python', '/home/oai/skills/slides/container_tools/render_slides.py', str(pptx_out)], check=True)
subprocess.run(['python', '/home/oai/skills/slides/container_tools/create_montage.py', '--input_dir', str(out / 'Lesson2_Controlled_Full_v0_5'), '--output_file', str(out / 'montage_full_v0_5.png')], check=True)
subprocess.run(['python', '/home/oai/skills/slides/container_tools/slides_test.py', str(pptx_out)], check=True)
